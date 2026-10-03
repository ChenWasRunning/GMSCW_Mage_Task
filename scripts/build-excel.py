import json,pathlib
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
root=pathlib.Path(__file__).resolve().parents[1];tasks=json.loads((root/'dist/tasks.json').read_text())
wb=Workbook();ws=wb.active;ws.title='完整流程'
ws.append(['完成','等级（预计）','任务与地图','目标与注意事项','材料获得／消耗'])
fmt=lambda m:'、'.join(f'{k} ×{v}' for k,v in m.items())
for t in tasks:
 ws.append(['☐',t['level'],f"{t['id']} · {t['phase']}\n{t['title']}\n{t['place']}",t['goal']+ ('\n注意：'+t['note'] if t['note'] else '')+(f"\n任务经验：{t['xp']}" if t['xp'] else ''),('获得／收齐：'+fmt(t['gain'])+'\n' if t['gain'] else '')+('消耗：'+fmt(t['cost']) if t['cost'] else '无材料消耗')])
ws.column_dimensions['A'].width=7;ws.column_dimensions['B'].width=14;ws.column_dimensions['C'].width=36;ws.column_dimensions['D'].width=55;ws.column_dimensions['E'].width=32
validation=DataValidation(type='list',formula1='"☐,☑"');ws.add_data_validation(validation);validation.add(f'A2:A{ws.max_row}')
ws.conditional_formatting.add(f'A2:E{ws.max_row}',FormulaRule(formula=['$A2="☑"'],font=Font(color='999999'),fill=PatternFill('solid',fgColor='EEEEEE')))
ws.auto_filter.ref=ws.dimensions
for row in range(2,ws.max_row+1):ws.row_dimensions[row].height=max(92,min(190,22*max(len(str(ws.cell(row,c).value))//(25 if c==3 else 38)+str(ws.cell(row,c).value).count('\n') for c in [3,4,5])))
ledger=wb.create_sheet('逐步累计背包');ledger.append(['步骤','任务','材料','此前数量','获得','消耗','完成后数量']);bag={}
for t in tasks:
 before=bag.copy()
 for k,n in t['gain'].items():bag[k]=bag.get(k,0)+n
 for k,n in t['cost'].items():bag[k]=bag.get(k,0)-n
 assert all(n>=0 for n in bag.values()),(t['id'],bag)
 for k in sorted(set(k for k,n in bag.items() if n>0)|set(t['cost'])):
  ledger.append([t['id'],t['title'],k,before.get(k,0),t['gain'].get(k,0),t['cost'].get(k,0),bag.get(k,0)])
for col,width in zip('ABCDEFG',[9,32,27,12,10,10,15]):ledger.column_dimensions[col].width=width
ledger.auto_filter.ref=ledger.dimensions
notes=wb.create_sheet('使用说明');notes.append(['说明','内容'])
for a,b in [('网页版','网页提供真实 checkbox、完成灰显、滑动隐藏、实时材料账本；进度仅存当前浏览器。'),('Excel','此表是完整路线计划快照。完成栏可选择 ☑ 并自动灰显；累计背包工作表按顺序全完成计算，不随 Excel 勾选重算。'),('数量规则','材料获得累加、提交扣除。未给数量的额外掉落、卖出、药品使用和金币余额不推测。'),('收集步骤','完成某一步表示该步列出的材料目标已经收齐；提前拾取的额外掉落可用于补足后续目标，不要再重复多刷。'),('等级','范围与经验百分比取自攻略，早期范围为路线估计。32+ 实际等级随额外练级变化。'),('来源','奇怪小鸭，2026-09-30；https://www.bilibili.com/video/BV1gQad6ZEJV/；用户提供全文转写，经典世界二测。'),('素材','MapleStory.io GMS v83 图标仅作历史图示。版权属于 NEXON。'),('纠错','黑猪→黑珠；倒伏／道服→道符；次蘑菇→刺蘑菇；人物异写统一。未展开的其他职业专属线和可选支线不扩写。'),('装备','法师职业鞋原文未给款式，待游戏核对；桑拿服男性蓝色／女性红色。'),('炼金','学徒母矿要求与炼金制作共计扣除一枚母矿，原文未称交学徒时另消耗一枚。'),('回城卷','法师写作业任务的回城卷奖励参照同一原文弓手段同名任务；法师段未重复讲奖励。')]:notes.append([a,b])
notes.column_dimensions['A'].width=16;notes.column_dimensions['B'].width=92
for sh in wb:
 sh.freeze_panes='C2' if sh==ws else 'A2';sh.sheet_view.showGridLines=False
 for cell in sh[1]:cell.fill=PatternFill('solid',fgColor='285647');cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');cell.alignment=Alignment(vertical='center',wrap_text=True)
 sh.row_dimensions[1].height=28
 for row in sh.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name='Microsoft YaHei',size=10,color='344A3B');cell.alignment=Alignment(vertical='top',wrap_text=True);cell.fill=PatternFill('solid',fgColor='F1F4EB' if cell.row%2==0 else 'FFFFFF')
 sh.sheet_properties.pageSetUpPr.fitToPage=True;sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A4;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0;sh.print_title_rows='1:1'
for r in range(2,notes.max_row+1):notes.row_dimensions[r].height=48
path=root/'dist/GMSCW_Mage_Task.xlsx';wb.save(path)
check=load_workbook(path);assert check['完整流程'].max_row==93
print('XLSX verified:',len(tasks),'tasks;',ledger.max_row-1,'inventory snapshot rows; final inventory:',bag)
